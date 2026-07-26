from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import re
import shutil
import tarfile
import zipfile
from collections import Counter, defaultdict
from dataclasses import dataclass
from pathlib import Path, PurePosixPath

from docx import Document
from pypdf import PdfReader


TEXT_EXTS = {".txt", ".md", ".tsv", ".csv", ".prov"}
AUDIT_EXTS = TEXT_EXTS | {".docx", ".pdf"}
VERSION_RE = re.compile(r"(?i)(?:^|[^a-z0-9])v(?:ersion)?[_ .-]?(\d+)(?:[_.](\d+))?(?:[_.](\d+))?")
PLACEHOLDER_RE = re.compile(r"(?i)(TODO|TBD|FIXME|XXX|\[встав|\{\{|<[^>]{1,40}>|заполнить|уточнить)")
NORMATIVE_RE = re.compile(r"(?i)\b(должен|должна|должны|необходимо|обязан|запрещено|следует|рекомендуется|может)\b")
CLAIM_RE = re.compile(r"(?i)\b(доказ|гарант|всегда|никогда|существенно|эффектив|надежн|валид|универсал)\w*")
REFERENCE_RE = re.compile(r"(?i)(см\.|раздел|приложен|таблиц|рисунк|источник|ГОСТ|ISO|doi|https?://)")


def sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def repair_name(name: str) -> str:
    # ZIPs created on another platform often expose UTF-8 bytes decoded as CP437.
    try:
        candidate = name.encode("cp437").decode("utf-8")
        if sum(ch in candidate for ch in "абвгдежзийклмнопрстуфхцчшщыэюяАБВГДЕЖЗИЙКЛМНОПРСТУФХЦЧШЩЫЭЮЯ") > 0:
            return candidate
    except (UnicodeEncodeError, UnicodeDecodeError):
        pass
    return name


def safe_slug(value: str) -> str:
    value = repair_name(value).replace("\\", "/")
    value = re.sub(r"[^0-9A-Za-zА-Яа-яЁё._/-]+", "_", value)
    return value.strip("/._") or "unnamed"


def decode_text(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-8", "cp1251", "cp866"):
        try:
            return repair_text(data.decode(enc))
        except UnicodeDecodeError:
            continue
    return repair_text(data.decode("utf-8", errors="replace"))


def repair_text(text: str) -> str:
    """Repair common UTF-8-as-Windows-1251 mojibake, only when score improves."""
    def badness(value: str) -> int:
        return sum(value.count(x) for x in ("Р", "С", "вЂ", "в„", "Р°", "СЃ", "С‚"))
    current = text
    for _ in range(2):
        try:
            candidate = current.encode("cp1251").decode("utf-8")
        except (UnicodeEncodeError, UnicodeDecodeError):
            break
        if badness(candidate) >= badness(current):
            break
        current = candidate
    return current


def extract_docx(data: bytes) -> list[tuple[str, str]]:
    doc = Document(io.BytesIO(data))
    out: list[tuple[str, str]] = []
    for idx, p in enumerate(doc.paragraphs, 1):
        text = " ".join(p.text.split())
        if text:
            out.append((f"p{idx}", text))
    for t_idx, table in enumerate(doc.tables, 1):
        for r_idx, row in enumerate(table.rows, 1):
            cells = [" ".join(cell.text.split()) for cell in row.cells]
            if any(cells):
                out.append((f"t{t_idx}r{r_idx}", " | ".join(cells)))
    return out


def extract_pdf(data: bytes) -> list[tuple[str, str]]:
    reader = PdfReader(io.BytesIO(data))
    out: list[tuple[str, str]] = []
    for page_idx, page in enumerate(reader.pages, 1):
        text = page.extract_text() or ""
        for line_idx, line in enumerate(text.splitlines(), 1):
            line = " ".join(line.split())
            if line:
                out.append((f"page{page_idx}l{line_idx}", line))
    return out


def extract_units(name: str, data: bytes) -> tuple[list[tuple[str, str]], str | None]:
    ext = PurePosixPath(name).suffix.lower()
    try:
        if ext == ".docx":
            return [(loc, repair_text(text)) for loc, text in extract_docx(data)], None
        if ext == ".pdf":
            return [(loc, repair_text(text)) for loc, text in extract_pdf(data)], None
        if ext in TEXT_EXTS:
            return [(f"l{i}", line.rstrip()) for i, line in enumerate(decode_text(data).splitlines(), 1) if line.strip()], None
    except Exception as exc:  # retain failure as audit evidence
        return [], f"{type(exc).__name__}: {exc}"
    return [], None


@dataclass
class Item:
    source_archive: str
    container_chain: str
    path: str
    repaired_path: str
    size: int
    sha256: str
    extension: str
    is_macos_metadata: bool
    depth: int


def iter_container(data: bytes, label: str, source: str, depth: int = 0):
    if depth > 5:
        return
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            for info in zf.infolist():
                if info.is_dir():
                    continue
                payload = zf.read(info)
                name = info.filename
                yield label, name, payload, depth
                ext = PurePosixPath(name).suffix.lower()
                if ext == ".zip" and len(payload) > 0:
                    yield from iter_container(payload, f"{label}!{repair_name(name)}", source, depth + 1)
                elif ext == ".tar" and len(payload) > 0:
                    try:
                        with tarfile.open(fileobj=io.BytesIO(payload)) as tf:
                            for member in tf.getmembers():
                                if member.isfile():
                                    f = tf.extractfile(member)
                                    if f:
                                        yield f"{label}!{repair_name(name)}", member.name, f.read(), depth + 1
                    except tarfile.TarError:
                        pass
    except zipfile.BadZipFile:
        return


def classify_line(text: str) -> list[str]:
    tags: list[str] = []
    if PLACEHOLDER_RE.search(text): tags.append("placeholder")
    if NORMATIVE_RE.search(text): tags.append("normative_language")
    if CLAIM_RE.search(text): tags.append("strong_claim")
    if REFERENCE_RE.search(text): tags.append("reference_or_citation")
    if len(text) > 500: tags.append("overlong_unit")
    if re.search(r"[!?]{2,}|\.{4,}", text): tags.append("punctuation_anomaly")
    if re.search(r"(?i)\b(пример|иллюстратив|условн)\w*\b", text): tags.append("illustrative_marker")
    return tags


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("source", type=Path)
    ap.add_argument("output", type=Path)
    args = ap.parse_args()
    out = args.output
    evidence = out / "evidence"
    extracted = out / "extracted_text"
    corpus = out / "corpus_unique"
    for p in (evidence, extracted, corpus): p.mkdir(parents=True, exist_ok=True)

    items: list[Item] = []
    by_hash: dict[str, list[dict]] = defaultdict(list)
    text_by_hash: dict[str, list[tuple[str, str]]] = {}
    error_by_hash: dict[str, str] = {}
    payload_by_hash: dict[str, tuple[str, bytes]] = {}
    source_zips = sorted(args.source.glob("*.zip"))

    for zp in source_zips:
        root_data = zp.read_bytes()
        for chain, raw_name, data, depth in iter_container(root_data, zp.name, zp.name):
            repaired = repair_name(raw_name)
            digest = sha256(data)
            ext = PurePosixPath(repaired).suffix.lower()
            macos = raw_name.startswith("__MACOSX/") or "/._" in raw_name or PurePosixPath(raw_name).name.startswith("._")
            item = Item(zp.name, chain, raw_name, repaired, len(data), digest, ext, macos, depth)
            items.append(item)
            by_hash[digest].append(item.__dict__)
            if not macos and ext in AUDIT_EXTS and digest not in text_by_hash:
                units, err = extract_units(repaired, data)
                text_by_hash[digest] = units
                if err: error_by_hash[digest] = err
                payload_by_hash[digest] = (repaired, data)

    # Machine-readable complete inventory.
    fields = list(Item.__annotations__)
    with (evidence / "inventory.csv").open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=fields); w.writeheader()
        for item in items: w.writerow(item.__dict__)
    (evidence / "duplicates.json").write_text(json.dumps({h: v for h, v in by_hash.items() if len(v) > 1}, ensure_ascii=False, indent=2), encoding="utf-8")

    # Exhaustive line/paragraph ledger and individual extracted texts.
    line_count = 0
    tag_counts: Counter[str] = Counter()
    duplicate_units: defaultdict[str, list[str]] = defaultdict(list)
    with (evidence / "line_audit.jsonl").open("w", encoding="utf-8") as ledger:
        for digest, units in sorted(text_by_hash.items()):
            name = payload_by_hash[digest][0]
            text_path = extracted / f"{digest[:16]}__{safe_slug(PurePosixPath(name).name)}.txt"
            with text_path.open("w", encoding="utf-8") as tf:
                for locator, text in units:
                    tags = classify_line(text)
                    tag_counts.update(tags)
                    unit_hash = sha256(re.sub(r"\s+", " ", text.strip().lower()).encode("utf-8"))
                    duplicate_units[unit_hash].append(f"{digest}:{locator}")
                    rec = {"file_sha256": digest, "source_name": name, "locator": locator, "text": text, "tags": tags, "normalized_text_sha256": unit_hash}
                    ledger.write(json.dumps(rec, ensure_ascii=False) + "\n")
                    tf.write(f"[{locator}] {text}\n")
                    line_count += 1

    cross_dupes = {h: refs for h, refs in duplicate_units.items() if len(refs) > 1}
    (evidence / "duplicate_units.json").write_text(json.dumps(cross_dupes, ensure_ascii=False, indent=2), encoding="utf-8")
    (evidence / "extraction_errors.json").write_text(json.dumps(error_by_hash, ensure_ascii=False, indent=2), encoding="utf-8")

    # Preserve one copy of each unique auditable artifact for reproducibility.
    for digest, (name, data) in payload_by_hash.items():
        target = corpus / f"{digest[:16]}__{safe_slug(PurePosixPath(name).name)}"
        target.write_bytes(data)

    ext_counts = Counter(i.extension or "[none]" for i in items if not i.is_macos_metadata)
    versions = Counter()
    for i in items:
        for match in VERSION_RE.finditer(i.repaired_path):
            versions[".".join(x or "0" for x in match.groups())] += 1
    summary = {
        "source_archives": len(source_zips), "container_entries_total": len(items),
        "non_macos_entries": sum(not i.is_macos_metadata for i in items),
        "unique_payloads": len(by_hash), "duplicate_hash_groups": sum(len(v) > 1 for v in by_hash.values()),
        "unique_auditable_files": len(text_by_hash), "audited_units": line_count,
        "extraction_errors": len(error_by_hash), "extensions": dict(ext_counts.most_common()),
        "version_mentions_in_paths": dict(versions.most_common()), "audit_tag_counts": dict(tag_counts),
    }
    (evidence / "summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()

